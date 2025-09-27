import tkinter as tk
from tkinter import ttk, messagebox, simpledialog
import pandas as pd
import os
import sqlite3
import subprocess
import time
import signal
import sys
import psutil
import csv
from datetime import datetime
from pathlib import Path

# Bu import'ları modüller mevcut değilse comment'e alın
try:
    from email_settings_window import EmailSettingsWindow
    from email_notification_system import email_system
    EMAIL_SYSTEM_AVAILABLE = True
except ImportError:
    EMAIL_SYSTEM_AVAILABLE = False
    print("E-posta sistemi modülleri bulunamadı, e-posta özellikleri devre dışı.")

def get_resource_path(relative_path):
    """PyInstaller ile paketlenmiş dosyalar için doğru yolu bulur"""
    try:
        # PyInstaller ile paketlendiğinde, geçici klasörde çalışır
        base_path = sys._MEIPASS
    except Exception:
        # Normal Python çalıştırma durumu
        base_path = os.path.abspath(".")
    
    return os.path.join(base_path, relative_path)

def get_executable_path():
    """Ana çalıştırılabilir dosyanın yolunu bulur - PyInstaller uyumlu"""
    executable_names = ["detect_and_recognize.exe", "detect_and_recognize.py"]
    
    # Önce exe dosyasının yanında ara
    try:
        if getattr(sys, 'frozen', False):
            # PyInstaller ile paketlenmiş durumda
            base_dir = os.path.dirname(sys.executable)
        else:
            # Normal Python çalıştırma durumu
            base_dir = os.path.dirname(os.path.abspath(__file__))
            
        for exe_name in executable_names:
            exe_path = os.path.join(base_dir, exe_name)
            if os.path.exists(exe_path):
                return exe_path, exe_name
        
        # PyInstaller resource path'te ara
        for exe_name in executable_names:
            try:
                resource_path = get_resource_path(exe_name)
                if os.path.exists(resource_path):
                    return resource_path, exe_name
            except:
                pass
                
        # Çalışma dizininde ara
        current_dir = os.getcwd()
        for exe_name in executable_names:
            exe_path = os.path.join(current_dir, exe_name)
            if os.path.exists(exe_path):
                return exe_path, exe_name
        
        # Hiçbir executable bulunamadıysa hata ver
        raise FileNotFoundError(
            "detect_and_recognize dosyası bulunamadı! Lütfen aşağıdaki konumlardan birinde dosyayı bulundurun:\n"
            f"- {base_dir}/detect_and_recognize.exe\n"
            f"- {base_dir}/detect_and_recognize.py\n"
            f"- {current_dir}/detect_and_recognize.exe\n"
            f"- {current_dir}/detect_and_recognize.py"
        )
    except Exception as e:
        print(f"Executable path bulma hatası: {e}")
        raise

def get_csv_path():
    """Her zaman çalışma dizininde CSV dosyası oluşturur/kullanır"""
    # EXE çalışıyorsa sys.executable'ın bulunduğu dizin, değilse current working dir
    if getattr(sys, 'frozen', False):
        base_dir = os.path.dirname(sys.executable)
    else:
        base_dir = os.getcwd()

    csv_path = os.path.join(base_dir, "detected_plates.csv")

    # CSV yoksa başlık satırıyla oluştur
    if not os.path.exists(csv_path):
        try:
            with open(csv_path, "w", newline='', encoding='utf-8') as csv_file:
                csv_writer = csv.writer(csv_file)
                csv_writer.writerow(["timestamp", "plate_number", "confidence"])
            print(f"Yeni CSV dosyası oluşturuldu: {csv_path}")
        except Exception as e:
            print(f"CSV dosyası oluşturulamadı: {e}")

    return csv_path


class DummyEmailSystem:
    """E-posta sistemi mevcut değilse kullanılacak sahte sınıf"""
    def is_configured(self):
        return False
    
    def send_test_email(self):
        return False, "E-posta sistemi yapılandırılmamış"
    
    @property
    def config(self):
        return {}

class PlateRecognitionApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Araç Plaka Tanıma Sistemi - Pro v2.0")
        self.root.geometry("1500x900")
        
        # E-posta sistemi kontrolü
        if EMAIL_SYSTEM_AVAILABLE:
            self.email_system = email_system
        else:
            self.email_system = DummyEmailSystem()
        
        # Plaka tanıma süreci değişkeni
        self.detection_process = None
        self.executable_path = None
        self.executable_name = None
        
        # Executable dosyasını bul ve kaydet
        try:
            self.executable_path, self.executable_name = get_executable_path()
            print(f"Executable bulundu: {self.executable_path}")
        except FileNotFoundError as e:
            print(f"Executable bulunamadı: {e}")
            # Hata durumunda None olarak bırak, start_detection'da tekrar kontrol edilecek
        
        # CSV path'ini al
        self.csv_path = get_csv_path()
        print(f"CSV yolu: {self.csv_path}")
        
        self.setup_ui()
        
        # CSV dosyası değişikliği izleme için değişken
        self.last_modified_time = None
        
        # Tüm verileri saklamak için değişken (arama için)
        self.all_plate_data = []
        
        # İlk başta CSV verilerini yükle
        self.load_csv_data()
        
        # Periyodik CSV kontrolü başlat
        self.check_csv_updates()
        
        # Uygulama kapatıldığında process'i durdur
        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)
    
    def setup_ui(self):
        """Kullanıcı arayüzünü kurar"""
        # Ana frame
        self.main_frame = tk.Frame(self.root)
        self.main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # Sol taraf - Durum ve Kontrol (genişletilmiş)
        self.left_frame = tk.Frame(self.main_frame, width=600)
        self.left_frame.pack(side=tk.LEFT, fill=tk.BOTH, padx=5, pady=5)
        self.left_frame.pack_propagate(False)  # Boyut sabitleme
        
        # Başlık
        self.title_label = tk.Label(self.left_frame, text="🚗 Plaka Tanıma Sistemi", 
                                   font=("Arial", 18, "bold"), fg="darkblue")
        self.title_label.pack(pady=10)
        
        # Durum etiketi (daha büyük)
        self.status_var = tk.StringVar(value="Durum: Beklemede")
        self.status_label = tk.Label(self.left_frame, textvariable=self.status_var,
                                   font=("Arial", 16, "bold"), pady=15, fg="darkgreen")
        self.status_label.pack(pady=15)
        
        # Sistem bilgileri frame'i - scrollbar ile
        self.system_frame = tk.LabelFrame(self.left_frame, text="Sistem Bilgileri", 
                                         font=("Arial", 12, "bold"), padx=10, pady=5)
        self.system_frame.pack(fill=tk.BOTH, expand=True, pady=10)
        
        # Sistem bilgileri için scrollable text widget
        self.system_text_frame = tk.Frame(self.system_frame)
        self.system_text_frame.pack(fill=tk.BOTH, expand=True)
        
        # Text widget ve scrollbar
        self.system_text = tk.Text(self.system_text_frame, height=8, width=50, 
                                  font=("Consolas", 10), wrap=tk.WORD, 
                                  bg="#f0f0f0", fg="black", state=tk.DISABLED)
        self.system_scrollbar = ttk.Scrollbar(self.system_text_frame, orient="vertical", 
                                             command=self.system_text.yview)
        
        self.system_text.configure(yscrollcommand=self.system_scrollbar.set)
        self.system_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.system_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.update_system_info()
        
        # E-posta durumu frame'i
        self.email_frame = tk.LabelFrame(self.left_frame, text="E-posta Bildirimleri", 
                                        font=("Arial", 12, "bold"), padx=10, pady=10)
        self.email_frame.pack(fill=tk.X, pady=10)
        
        self.email_status_var = tk.StringVar()
        self.email_status_label = tk.Label(self.email_frame, textvariable=self.email_status_var,
                                          font=("Arial", 11), pady=5, justify=tk.LEFT)
        self.email_status_label.pack()
        self.update_email_status()
        
        # E-posta butonları
        self.email_button_frame = tk.Frame(self.email_frame)
        self.email_button_frame.pack(fill=tk.X, pady=5)
        
        self.email_settings_button = tk.Button(self.email_button_frame, text="⚙️ E-posta Ayarları", 
                                             command=self.open_email_settings, bg="#2196F3", fg="white",
                                             font=("Arial", 10), padx=10, pady=5, width=18,
                                             state=tk.NORMAL if EMAIL_SYSTEM_AVAILABLE else tk.DISABLED)
        self.email_settings_button.pack(side=tk.LEFT, padx=(0, 5))
        
        self.test_email_button = tk.Button(self.email_button_frame, text="📧 Test Gönder", 
                                         command=self.send_test_email, bg="#FF9800", fg="white",
                                         font=("Arial", 10), padx=10, pady=5, width=18,
                                         state=tk.NORMAL if EMAIL_SYSTEM_AVAILABLE else tk.DISABLED)
        self.test_email_button.pack(side=tk.LEFT)
        
        # Ana kontrol butonları için frame
        self.main_button_frame = tk.Frame(self.left_frame)
        self.main_button_frame.pack(fill=tk.X, pady=20)
        
        self.start_button = tk.Button(self.main_button_frame, text="🚀 PLAKA TANIMA BAŞLAT", 
                                     command=self.start_detection, bg="#4CAF50", fg="white",
                                     font=("Arial", 14, "bold"), padx=20, pady=15,
                                     width=30, height=1)
        self.start_button.pack(pady=5)
        
        self.stop_button = tk.Button(self.main_button_frame, text="⏹️ PLAKA TANIMA DURDUR", 
                                    command=self.stop_detection, bg="#F44336", fg="white",
                                    font=("Arial", 14, "bold"), padx=20, pady=15,
                                    width=30, height=1, state=tk.DISABLED)
        self.stop_button.pack(pady=5)
        
        # İstatistikler frame - scrollbar ile
        self.stats_main_frame = tk.LabelFrame(self.left_frame, text="📊 İstatistikler", 
                                             font=("Arial", 12, "bold"), padx=10, pady=10)
        self.stats_main_frame.pack(fill=tk.BOTH, expand=True, pady=10)
        
        # İstatistikler için scrollable frame
        self.setup_scrollable_stats_frame()
        
        # Sağ taraf - Plaka listesi (genişletilmiş)
        self.right_frame = tk.Frame(self.main_frame)
        self.right_frame.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        # Arama özelliği için frame
        self.search_frame = tk.Frame(self.right_frame)
        self.search_frame.pack(fill=tk.X, pady=(10, 5))
        
        # Arama etiketi
        self.search_label = tk.Label(self.search_frame, text="🔍 Plaka Ara:", font=("Arial", 12, "bold"))
        self.search_label.pack(side=tk.LEFT, padx=(0, 10))
        
        # Arama giriş kutusu
        self.search_var = tk.StringVar()
        self.search_entry = tk.Entry(self.search_frame, textvariable=self.search_var, 
                                    font=("Arial", 12), width=15)
        self.search_entry.pack(side=tk.LEFT, padx=(0, 10))
        self.search_entry.bind('<Return>', lambda e: self.search_plates())
        
        # Arama butonu
        self.search_button = tk.Button(self.search_frame, text="🔍 Ara", 
                                      command=self.search_plates, bg="#2196F3", fg="white",
                                      font=("Arial", 10), padx=10, pady=5)
        self.search_button.pack(side=tk.LEFT, padx=(0, 10))
        
        # Filtreyi temizle butonu
        self.clear_filter_button = tk.Button(self.search_frame, text="🗑️ Filtreyi Temizle", 
                                           command=self.clear_filter, bg="#FF9800", fg="white",
                                           font=("Arial", 10), padx=10, pady=5)
        self.clear_filter_button.pack(side=tk.LEFT)
        
        # Liste başlığı
        self.header_label = tk.Label(self.right_frame, text="📋 Tespit Edilen Plakalar", 
                                    font=("Arial", 16, "bold"), fg="darkblue")
        self.header_label.pack(pady=(10, 10))
        
        # Treeview frame (scrollbar için)
        self.tree_frame = tk.Frame(self.right_frame)
        self.tree_frame.pack(fill=tk.BOTH, expand=True)
        
        # Treeview oluşturma
        self.columns = ("tarih_saat", "plaka", "dogruluk")
        self.plate_tree = ttk.Treeview(self.tree_frame, columns=self.columns, show="headings")
        
        # Kolonları tanımla
        self.plate_tree.heading("tarih_saat", text="📅 Tarih ve Saat")
        self.plate_tree.heading("plaka", text="🚗 Plaka No")
        self.plate_tree.heading("dogruluk", text="✅ Doğruluk Oranı")
        
        # Kolon genişlikleri
        self.plate_tree.column("tarih_saat", width=200, anchor="center")
        self.plate_tree.column("plaka", width=150, anchor="center")
        self.plate_tree.column("dogruluk", width=150, anchor="center")
        
        # Scrollbar'ları ekle (hem dikey hem yatay)
        self.tree_v_scroll = ttk.Scrollbar(self.tree_frame, orient="vertical", 
                                          command=self.plate_tree.yview)
        self.tree_h_scroll = ttk.Scrollbar(self.tree_frame, orient="horizontal", 
                                          command=self.plate_tree.xview)
        
        self.plate_tree.configure(yscrollcommand=self.tree_v_scroll.set,
                                 xscrollcommand=self.tree_h_scroll.set)
        
        # Treeview ve scrollbar'ları yerleştir
        self.plate_tree.grid(row=0, column=0, sticky="nsew")
        self.tree_v_scroll.grid(row=0, column=1, sticky="ns")
        self.tree_h_scroll.grid(row=1, column=0, sticky="ew")
        
        # Grid yapılandırması
        self.tree_frame.grid_rowconfigure(0, weight=1)
        self.tree_frame.grid_columnconfigure(0, weight=1)
        
        # Alt butonlar için frame
        self.bottom_button_frame = tk.Frame(self.right_frame)
        self.bottom_button_frame.pack(fill=tk.X, pady=15)
        
        # Kaydet butonu
        self.save_button = tk.Button(self.bottom_button_frame, text="💾 Veritabanına Kaydet", 
                                    command=self.save_to_database, bg="#2196F3", fg="white",
                                    font=("Arial", 12), padx=15, pady=8)
        self.save_button.pack(side=tk.LEFT, padx=(0, 10))
        
        # Yenile butonu
        self.refresh_button = tk.Button(self.bottom_button_frame, text="🔄 Listeyi Yenile", 
                                      command=self.load_csv_data, bg="#4CAF50", fg="white",
                                      font=("Arial", 12), padx=15, pady=8)
        self.refresh_button.pack(side=tk.LEFT, padx=(0, 10))
        
        # Excel'e aktar butonu
        self.export_button = tk.Button(self.bottom_button_frame, text="📊 Excel'e Aktar", 
                                     command=self.export_to_excel, bg="#009688", fg="white",
                                     font=("Arial", 12), padx=15, pady=8)
        self.export_button.pack(side=tk.LEFT, padx=(0, 10))
        
        # Listeyi temizle butonu
        self.clear_button = tk.Button(self.bottom_button_frame, text="🗑️ Listeyi Temizle", 
                                    command=self.clear_plate_list, bg="#F44336", fg="white",
                                    font=("Arial", 12), padx=15, pady=8)
        self.clear_button.pack(side=tk.LEFT)
    
    def setup_scrollable_stats_frame(self):
        """İstatistikler için scrollable frame oluşturur"""
        # Canvas ve scrollbar için frame
        self.stats_canvas_frame = tk.Frame(self.stats_main_frame)
        self.stats_canvas_frame.pack(fill=tk.BOTH, expand=True)
        
        # Canvas oluştur
        self.stats_canvas = tk.Canvas(self.stats_canvas_frame, height=200)
        self.stats_v_scrollbar = ttk.Scrollbar(self.stats_canvas_frame, orient="vertical", 
                                              command=self.stats_canvas.yview)
        
        # Scrollable frame
        self.stats_scrollable_frame = tk.Frame(self.stats_canvas)
        
        self.stats_scrollable_frame.bind(
    "<Configure>",
    lambda e: self.stats_canvas.configure(scrollregion=self.stats_canvas.bbox("all"))
)
        
        self.stats_canvas.create_window((0, 0), window=self.stats_scrollable_frame, anchor="nw")
        self.stats_canvas.configure(yscrollcommand=self.stats_v_scrollbar.set)
        
        # Canvas ve scrollbar'ı yerleştir
        self.stats_canvas.pack(side="left", fill="both", expand=True)
        self.stats_v_scrollbar.pack(side="right", fill="y")
        
        # Mouse wheel binding
        def _on_mousewheel(event):
            self.stats_canvas.yview_scroll(int(-1*(event.delta/120)), "units")
        self.stats_canvas.bind("<MouseWheel>", _on_mousewheel)
        
        # İstatistik etiketleri
        self.setup_stats_labels()
    
    def setup_stats_labels(self):
        """İstatistik etiketlerini oluşturur"""
        # Temel istatistikler
        self.total_plates_var = tk.StringVar(value="📊 Toplam Plaka: 0")
        self.total_plates_label = tk.Label(self.stats_scrollable_frame, textvariable=self.total_plates_var,
                                          font=("Arial", 12, "bold"), anchor=tk.W, fg="darkblue")
        self.total_plates_label.pack(fill=tk.X, pady=3)
        
        self.today_plates_var = tk.StringVar(value="📅 Bugün: 0")
        self.today_plates_label = tk.Label(self.stats_scrollable_frame, textvariable=self.today_plates_var,
                                          font=("Arial", 11), anchor=tk.W, fg="darkgreen")
        self.today_plates_label.pack(fill=tk.X, pady=2)
        
        self.last_plate_var = tk.StringVar(value="🚗 Son Plaka: -")
        self.last_plate_label = tk.Label(self.stats_scrollable_frame, textvariable=self.last_plate_var,
                                        font=("Arial", 11), anchor=tk.W, fg="darkorange")
        self.last_plate_label.pack(fill=tk.X, pady=2)
        
        # Ayırıcı çizgi
        separator1 = ttk.Separator(self.stats_scrollable_frame, orient='horizontal')
        separator1.pack(fill=tk.X, pady=5)
        
        # Ek istatistikler
        self.unique_plates_var = tk.StringVar(value="🔢 Benzersiz Plaka: 0")
        self.unique_plates_label = tk.Label(self.stats_scrollable_frame, textvariable=self.unique_plates_var,
                                          font=("Arial", 11), anchor=tk.W, fg="purple")
        self.unique_plates_label.pack(fill=tk.X, pady=2)
        
        self.avg_confidence_var = tk.StringVar(value="📈 Ortalama Doğruluk: 0%")
        self.avg_confidence_label = tk.Label(self.stats_scrollable_frame, textvariable=self.avg_confidence_var,
                                           font=("Arial", 11), anchor=tk.W, fg="darkred")
        self.avg_confidence_label.pack(fill=tk.X, pady=2)
        
        self.highest_confidence_var = tk.StringVar(value="🏆 En Yüksek Doğruluk: 0%")
        self.highest_confidence_label = tk.Label(self.stats_scrollable_frame, textvariable=self.highest_confidence_var,
                                                font=("Arial", 11), anchor=tk.W, fg="darkmagenta")
        self.highest_confidence_label.pack(fill=tk.X, pady=2)
        
        # Ayırıcı çizgi
        separator2 = ttk.Separator(self.stats_scrollable_frame, orient='horizontal')
        separator2.pack(fill=tk.X, pady=5)
        
        # Sistem durumu
        self.csv_status_var = tk.StringVar(value="📄 CSV Durumu: Kontrol ediliyor...")
        self.csv_status_label = tk.Label(self.stats_scrollable_frame, textvariable=self.csv_status_var,
                                        font=("Arial", 10), anchor=tk.W, fg="darkcyan")
        self.csv_status_label.pack(fill=tk.X, pady=2)
        
        self.system_uptime_var = tk.StringVar(value="⏱️ Çalışma Süresi: 0 dk")
        self.system_uptime_label = tk.Label(self.stats_scrollable_frame, textvariable=self.system_uptime_var,
                                          font=("Arial", 10), anchor=tk.W, fg="gray")
        self.system_uptime_label.pack(fill=tk.X, pady=2)
        
        # Başlangıç zamanını kaydet
        self.start_time = time.time()
    
    def update_system_info(self):
        """Sistem bilgilerini güncelle"""
        try:
            info_lines = []
            
            if getattr(sys, 'frozen', False):
                mode = "🔧 Mod: EXE (PyInstaller)"
                base_path = os.path.dirname(sys.executable)
            else:
                mode = "🔧 Mod: Python Geliştirme"
                base_path = os.path.dirname(os.path.abspath(__file__))
            
            info_lines.append(mode)
            
            executable_status = "✅ Executable: Bulundu" if self.executable_path else "❌ Executable: Bulunamadı"
            info_lines.append(executable_status)
            
            if self.executable_path:
                info_lines.append(f"📂 Dosya: {os.path.basename(self.executable_path)}")
            
            info_lines.append(f"📁 Çalışma Dizini: {os.getcwd()}")
            info_lines.append(f"📄 CSV Yolu: {os.path.basename(self.csv_path)}")
            
            # CSV dosyası durumu
            if os.path.exists(self.csv_path):
                csv_size = os.path.getsize(self.csv_path)
                info_lines.append(f"📊 CSV Boyutu: {csv_size} byte")
                
                # CSV yazma yetkisi
                csv_writable = os.access(self.csv_path, os.W_OK)
                write_status = "✅ Yazılabilir" if csv_writable else "❌ Yazılamaz"
                info_lines.append(f"✏️ CSV Yazma: {write_status}")
            else:
                info_lines.append("❌ CSV: Dosya mevcut değil")
            
            # Sistem kaynak kullanımı
            try:
                import psutil
                process = psutil.Process()
                memory_mb = process.memory_info().rss / 1024 / 1024
                cpu_percent = process.cpu_percent()
                info_lines.append(f"🖥️ Bellek: {memory_mb:.1f} MB")
                info_lines.append(f"💻 CPU: {cpu_percent:.1f}%")
            except:
                pass
            
            # E-posta sistem durumu
            if EMAIL_SYSTEM_AVAILABLE:
                if self.email_system.is_configured():
                    info_lines.append("📧 E-posta: Yapılandırıldı")
                else:
                    info_lines.append("📧 E-posta: Yapılandırılmamış")
            else:
                info_lines.append("📧 E-posta: Modül mevcut değil")
            
            # Text widget'ı güncelle
            self.system_text.config(state=tk.NORMAL)
            self.system_text.delete(1.0, tk.END)
            self.system_text.insert(tk.END, "\n".join(info_lines))
            self.system_text.config(state=tk.DISABLED)
            
        except Exception as e:
            self.system_text.config(state=tk.NORMAL)
            self.system_text.delete(1.0, tk.END)
            self.system_text.insert(tk.END, f"❌ Sistem bilgisi hatası:\n{e}")
            self.system_text.config(state=tk.DISABLED)
    
    def update_email_status(self):
        """E-posta durumunu güncelle"""
        if not EMAIL_SYSTEM_AVAILABLE:
            self.email_status_var.set("❌ E-posta Sistemi: Modül bulunamadı\nE-posta bildirimleri kullanılamaz.")
            self.email_status_label.config(fg="gray")
            return
            
        if self.email_system.is_configured():
            config = self.email_system.config
            recipient = config.get("recipient_email", "Bilinmiyor")
            if config.get("notifications_enabled", False):
                self.email_status_var.set(f"✅ E-posta Bildirimi: AKTİF\n📧 Alıcı: {recipient}")
                self.email_status_label.config(fg="green")
            else:
                self.email_status_var.set(f"⚠️ E-posta Bildirimi: Pasif\n📧 Alıcı: {recipient}")
                self.email_status_label.config(fg="orange")
        else:
            self.email_status_var.set("❌ E-posta Bildirimi: Yapılandırılmamış\nAyarları yapmak için butona tıklayın.")
            self.email_status_label.config(fg="red")
    
    def update_statistics(self):
        """İstatistikleri güncelle"""
        try:
            # Temel istatistikler
            total_count = len(self.all_plate_data)
            self.total_plates_var.set(f"📊 Toplam Plaka: {total_count}")
            
            # Bugünkü plakalar
            today = datetime.now().strftime("%Y-%m-%d")
            today_count = sum(1 for plate in self.all_plate_data 
                            if plate["tarih_saat"].startswith(today))
            self.today_plates_var.set(f"📅 Bugün: {today_count}")
            
            # Son plaka
            if self.all_plate_data:
                last_plate = self.all_plate_data[0]["plaka"]  # En yeni kayıt
                self.last_plate_var.set(f"🚗 Son Plaka: {last_plate}")
            else:
                self.last_plate_var.set("🚗 Son Plaka: -")
            
            # Benzersiz plakalar
            unique_plates = set(plate["plaka"] for plate in self.all_plate_data)
            self.unique_plates_var.set(f"🔢 Benzersiz Plaka: {len(unique_plates)}")
            
            # Ortalama doğruluk oranı
            if self.all_plate_data:
                confidences = []
                for plate in self.all_plate_data:
                    try:
                        conf_str = plate["dogruluk"].replace("%", "")
                        confidences.append(float(conf_str))
                    except:
                        continue
                
                if confidences:
                    avg_conf = sum(confidences) / len(confidences)
                    max_conf = max(confidences)
                    self.avg_confidence_var.set(f"📈 Ortalama Doğruluk: {avg_conf:.1f}%")
                    self.highest_confidence_var.set(f"🏆 En Yüksek Doğruluk: {max_conf:.1f}%")
                else:
                    self.avg_confidence_var.set("📈 Ortalama Doğruluk: N/A")
                    self.highest_confidence_var.set("🏆 En Yüksek Doğruluk: N/A")
            else:
                self.avg_confidence_var.set("📈 Ortalama Doğruluk: 0%")
                self.highest_confidence_var.set("🏆 En Yüksek Doğruluk: 0%")
            
            # CSV durumu
            if os.path.exists(self.csv_path):
                csv_writable = os.access(self.csv_path, os.W_OK)
                if csv_writable:
                    self.csv_status_var.set("📄 CSV Durumu: ✅ Yazılabilir")
                else:
                    self.csv_status_var.set("📄 CSV Durumu: ❌ Yazılamaz")
            else:
                self.csv_status_var.set("📄 CSV Durumu: ❌ Dosya mevcut değil")
            
            # Sistem çalışma süresi
            uptime_seconds = time.time() - self.start_time
            uptime_minutes = int(uptime_seconds / 60)
            uptime_hours = int(uptime_minutes / 60)
            remaining_minutes = uptime_minutes % 60
            
            if uptime_hours > 0:
                uptime_str = f"{uptime_hours}s {remaining_minutes}dk"
            else:
                uptime_str = f"{uptime_minutes}dk"
            
            self.system_uptime_var.set(f"⏱️ Çalışma Süresi: {uptime_str}")
                
        except Exception as e:
            print(f"İstatistik güncelleme hatası: {e}")
    
    def load_csv_data(self):
        """CSV dosyasından plaka verilerini yükler"""
        try:
            # CSV dosyası yoksa oluştur
            if not os.path.exists(self.csv_path):
                with open(self.csv_path, "w", newline='', encoding='utf-8') as csv_file:
                    csv_writer = csv.writer(csv_file)
                    csv_writer.writerow(["timestamp", "plate_number", "confidence"])
                print(f"Yeni CSV dosyası oluşturuldu: {self.csv_path}")
                self.all_plate_data = []
                self.display_plates([])
                self.update_statistics()
                return
            
            # CSV dosyasının son değiştirilme zamanını kaydet
            self.last_modified_time = os.path.getmtime(self.csv_path)
            
            print(f"CSV dosyası okunuyor: {self.csv_path}")
            
            # Önce dosyanın içeriğini kontrol et
            with open(self.csv_path, 'r', encoding='utf-8') as file:
                content = file.read().strip()
                if not content or content == "timestamp,plate_number,confidence":
                    print("CSV dosyası boş veya sadece header içeriyor.")
                    self.all_plate_data = []
                    self.display_plates([])
                    self.update_statistics()
                    return
            
            # CSV dosyasını oku
            self.all_plate_data = []
            
            with open(self.csv_path, 'r', encoding='utf-8') as file:
                csv_reader = csv.reader(file)
                header = next(csv_reader, None)  # Başlık satırını atla
                
                for i, row in enumerate(csv_reader):
                    if len(row) >= 3:
                        try:
                            timestamp = row[0].strip()
                            plate_number = row[1].strip()
                            confidence = float(row[2].strip())
                            
                            # Doğruluk değerini yüzde formatına çevir
                            if confidence <= 1.0:  # Eğer decimal formatında ise
                                dogruluk = f"{confidence * 100:.2f}%"
                            else:  # Zaten yüzde formatında ise
                                dogruluk = f"{confidence:.2f}%"
                            
                            self.all_plate_data.append({
                                "tarih_saat": timestamp,
                                "plaka": plate_number,
                                "dogruluk": dogruluk
                            })
                            
                        except Exception as e:
                            print(f"Satır {i+1} işlenirken hata: {e} - Satır: {row}")
            
            print(f"Toplam {len(self.all_plate_data)} kayıt yüklendi")
            
            # Verileri tarihe göre sırala (yeni kayıtlar üstte)
            self.all_plate_data.sort(key=lambda x: x["tarih_saat"], reverse=True)
            
            # Verileri listele
            self.display_plates(self.all_plate_data)
            
            # İstatistikleri güncelle
            self.update_statistics()
            
        except Exception as e:
            print(f"CSV dosyası yüklenirken hata oluştu: {e}")
            messagebox.showerror("Hata", f"CSV dosyası yüklenirken hata oluştu: {e}")
    
    def display_plates(self, plate_data):
        """Belirtilen plaka verilerini görüntüler"""
        # Önce listeyi temizle
        for item in self.plate_tree.get_children():
            self.plate_tree.delete(item)
        
        print(f"Görüntülenecek kayıt sayısı: {len(plate_data)}")
        
        # Verileri ekle
        for i, plate in enumerate(plate_data):
            try:
                # Alternatif satır renkleri için tag kullan
                tag = 'even' if i % 2 == 0 else 'odd'
                
                self.plate_tree.insert("", "end", values=(
                    plate["tarih_saat"],
                    plate["plaka"],
                    plate["dogruluk"]
                ), tags=(tag,))
                
            except Exception as e:
                print(f"Kayıt {i+1} eklenirken hata: {e}")
        
        # Alternatif satır renkleri
        self.plate_tree.tag_configure('odd', background='#f0f0f0')
        self.plate_tree.tag_configure('even', background='white')
    
    def search_plates(self):
        """Arama kutusundaki metne göre plakaları filtreler"""
        search_text = self.search_var.get().strip().upper()
        if not search_text:
            # Arama metni boşsa tüm listeyi göster
            self.display_plates(self.all_plate_data)
            return
        
        # Arama metni varsa filtreleme yap
        filtered_data = [plate for plate in self.all_plate_data 
                        if search_text in plate["plaka"].upper() or 
                           search_text in plate["tarih_saat"]]
        
        # Filtrelenmiş verileri göster
        self.display_plates(filtered_data)
        
        # Sonuç sayısını bildir
        if not filtered_data:
            messagebox.showinfo("Arama Sonucu", "🔍 Aramanızla eşleşen plaka bulunamadı.")
        else:
            messagebox.showinfo("Arama Sonucu", f"🔍 {len(filtered_data)} plaka bulundu.")
    
    def clear_filter(self):
        """Filtreyi temizler ve tüm plakaları gösterir"""
        self.search_var.set("")  # Arama kutusunu temizle
        self.display_plates(self.all_plate_data)  # Tüm verileri göster
    
    def export_to_excel(self):
        """Plaka verilerini Excel dosyasına aktarır"""
        try:
            if not self.all_plate_data:
                messagebox.showinfo("Bilgi", "📊 Aktarılacak veri bulunamadı.")
                return
            
            # Pandas DataFrame oluştur
            df = pd.DataFrame(self.all_plate_data)
            
            # Excel dosya adı (tarih ile)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            excel_filename = f"plaka_listesi_{timestamp}.xlsx"
            excel_path = os.path.join(os.getcwd(), excel_filename)
            
            # Excel'e yaz
            with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name='Plakalar', index=False)
                
                # Worksheet'i al ve formatla
                worksheet = writer.sheets['Plakalar']
                
                # Sütun genişliklerini ayarla
                worksheet.column_dimensions['A'].width = 20  # Tarih
                worksheet.column_dimensions['B'].width = 15  # Plaka
                worksheet.column_dimensions['C'].width = 15  # Doğruluk
                
                # Başlık satırını formatla
                from openpyxl.styles import Font, PatternFill
                header_font = Font(bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                
                for col in range(1, len(df.columns) + 1):
                    cell = worksheet.cell(row=1, column=col)
                    cell.font = header_font
                    cell.fill = header_fill
            
            messagebox.showinfo("Başarılı", f"📊 Veriler Excel dosyasına aktarıldı:\n{excel_path}")
            
        except ImportError:
            messagebox.showerror("Hata", "📊 Excel aktarımı için pandas ve openpyxl kütüphaneleri gerekli.\n\nKurulum:\npip install pandas openpyxl")
        except Exception as e:
            messagebox.showerror("Hata", f"📊 Excel aktarımı sırasında hata: {e}")
    
    def clear_plate_list(self):
        """Plaka listesini temizler (CSV ve arayüz)"""
        confirm = messagebox.askyesno("⚠️ Onay", 
                                    "Plaka listesi tamamen temizlenecek.\n"
                                    "Bu işlem geri alınamaz!\n\n"
                                    "Devam etmek istiyor musunuz?")
        
        if confirm:
            try:
                # CSV dosyasını başlık satırı ile yeniden oluştur
                with open(self.csv_path, 'w', newline='', encoding='utf-8') as file:
                    csv_writer = csv.writer(file)
                    csv_writer.writerow(["timestamp", "plate_number", "confidence"])
                
                # Arayüzdeki listeyi temizle
                for item in self.plate_tree.get_children():
                    self.plate_tree.delete(item)
                
                # Veri listesini temizle
                self.all_plate_data = []
                
                # İstatistikleri güncelle
                self.update_statistics()
                
                messagebox.showinfo("✅ Bilgi", "Plaka listesi başarıyla temizlendi.")
            except Exception as e:
                messagebox.showerror("❌ Hata", f"Liste temizlenirken hata oluştu: {e}")
    
    def check_csv_updates(self):
        """CSV dosyası değişikliklerini periyodik olarak kontrol eder"""
        try:
            if os.path.exists(self.csv_path):
                current_modified_time = os.path.getmtime(self.csv_path)
                if self.last_modified_time is None or current_modified_time > self.last_modified_time:
                    self.load_csv_data()
                    # İstatistikleri ve sistem bilgilerini güncelle
                    self.update_statistics()
                    self.update_system_info()
                    self.update_email_status()
        except Exception as e:
            print(f"CSV güncelleme kontrolü hatası: {e}")
        
        # Her 2 saniyede bir kontrol et
        self.root.after(2000, self.check_csv_updates)
    
    def open_email_settings(self):
        """E-posta ayarları penceresini aç"""
        if not EMAIL_SYSTEM_AVAILABLE:
            messagebox.showerror("❌ Hata", "E-posta sistemi modülleri bulunamadı!")
            return
            
        email_window = EmailSettingsWindow(self.root)
        
        # Pencere kapandığında e-posta durumunu güncelle
        def on_email_window_close():
            # Config'i yeniden yükle
            self.email_system.config = self.email_system.load_config()
            self.update_email_status()
        
        email_window.window.protocol("WM_DELETE_WINDOW", lambda: [
            email_window.window.destroy(),
            on_email_window_close()
        ])
    
    def send_test_email(self):
        """Test e-postası gönder"""
        if not EMAIL_SYSTEM_AVAILABLE:
            messagebox.showerror("❌ Hata", "E-posta sistemi modülleri bulunamadı!")
            return
            
        if not self.email_system.is_configured():
            messagebox.showerror("❌ Hata", 
                               "E-posta ayarları yapılandırılmamış!\nÖnce '⚙️ E-posta Ayarları' butonuna tıklayın.")
            return
        
        try:
            # Test e-postası gönder
            success, message = self.email_system.send_test_email()
            if success:
                messagebox.showinfo("✅ Başarılı", f"📧 {message}")
            else:
                messagebox.showerror("❌ Hata", f"📧 {message}")
        except Exception as e:
            messagebox.showerror("❌ Hata", f"Test e-postası gönderilirken hata: {e}")
    
    def start_detection(self):
        """Plaka tanıma sistemini başlatır"""
        if self.detection_process is None:
            try:
                # Executable path'ini tekrar kontrol et
                if not self.executable_path:
                    try:
                        self.executable_path, self.executable_name = get_executable_path()
                        self.update_system_info()
                    except FileNotFoundError as e:
                        messagebox.showerror("❌ Hata", str(e))
                        return
                
                # Çalışma dizinini executable'ın bulunduğu dizin olarak ayarla
                work_dir = os.path.dirname(self.executable_path)
                
                print(f"Başlatılacak dosya: {self.executable_path}")
                print(f"Çalışma dizini: {work_dir}")
                
                # Executable tipine göre çalıştırma yöntemi belirle
                if self.executable_name.endswith('.exe'):
                    # EXE dosyası ise doğrudan çalıştır
                    self.detection_process = subprocess.Popen([self.executable_path], 
                                                            cwd=work_dir,
                                                            creationflags=subprocess.CREATE_NO_WINDOW)
                else:
                    # Python dosyası ise python ile çalıştır
                    self.detection_process = subprocess.Popen(["python", self.executable_path], 
                                                            cwd=work_dir)
                
                # Butonların durumunu güncelle
                self.start_button.config(state=tk.DISABLED, bg="#cccccc")
                self.stop_button.config(state=tk.NORMAL, bg="#F44336")
                
                # Durum etiketini güncelle
                self.status_var.set("🚀 Durum: Plaka Tanıma AKTİF")
                self.status_label.config(fg="green")
                
                # E-posta durumu kontrolü
                email_status = ""
                if EMAIL_SYSTEM_AVAILABLE:
                    email_status = "📧 E-posta bildirimleri de aktif." if self.email_system.is_configured() else "⚠️ E-posta bildirimleri yapılandırılmamış."
                else:
                    email_status = "ℹ️ E-posta sistemi mevcut değil."
                
                messagebox.showinfo("✅ Başarılı", f"🚀 Plaka tanıma sistemi başlatıldı.\n{email_status}")
                
            except Exception as e:
                messagebox.showerror("❌ Hata", f"Plaka tanıma sistemi başlatılırken hata: {e}")
                # Hata durumunda butonları sıfırla
                self.start_button.config(state=tk.NORMAL, bg="#4CAF50")
                self.stop_button.config(state=tk.DISABLED, bg="#cccccc")
                self.status_var.set("❌ Durum: Hata - Başlatılamadı")
                self.status_label.config(fg="red")
    
    def stop_detection(self):
        """Plaka tanıma sistemini durdurur"""
        if self.detection_process is not None:
            try:
                # Prosesi nazikçe kapatmayı dene
                self.terminate_process()
                
                # Butonların durumunu güncelle
                self.start_button.config(state=tk.NORMAL, bg="#4CAF50")
                self.stop_button.config(state=tk.DISABLED, bg="#cccccc")
                
                # Durum etiketini güncelle
                self.status_var.set("⏹️ Durum: Beklemede")
                self.status_label.config(fg="darkorange")
                
                messagebox.showinfo("✅ Bilgi", "⏹️ Plaka tanıma sistemi durduruldu.")
            except Exception as e:
                messagebox.showerror("❌ Hata", f"Plaka tanıma sistemi durdurulurken hata: {e}")
    
    def terminate_process(self):
        """İşlemi ve tüm alt işlemlerini güvenli bir şekilde sonlandırır"""
        if self.detection_process is not None:
            try:
                # İşlem hala çalışıyor mu kontrol et
                if self.detection_process.poll() is None:
                    try:
                        process = psutil.Process(self.detection_process.pid)
                        
                        # Alt işlemleri de sonlandır
                        for child in process.children(recursive=True):
                            try:
                                child.terminate()
                            except:
                                pass
                        
                        # Ana işlemi sonlandır
                        self.detection_process.terminate()
                        
                        # İşlemin kapanmasını en fazla 5 saniye bekle
                        try:
                            self.detection_process.wait(timeout=5)
                        except subprocess.TimeoutExpired:
                            # Hala kapanmadıysa zorla kapat
                            print("İşlem zorla sonlandırılıyor...")
                            for child in process.children(recursive=True):
                                try:
                                    child.kill()
                                except:
                                    pass
                            self.detection_process.kill()
                            self.detection_process.wait()  # Kesin kapanana kadar bekle
                            
                    except psutil.NoSuchProcess:
                        # İşlem zaten kapatılmış
                        pass
                    except Exception as e:
                        print(f"Psutil ile sonlandırma hatası: {e}")
                        # Psutil başarısız olursa basit terminate dene
                        try:
                            self.detection_process.terminate()
                            self.detection_process.wait(timeout=3)
                        except:
                            self.detection_process.kill()
                            
            except Exception as e:
                print(f"İşlem sonlandırma hatası: {e}")
            
            # İşlem değişkenini sıfırla
            self.detection_process = None
    
    def save_to_database(self):
        """CSV dosyasındaki verileri veritabanına kaydeder"""
        try:
            if not self.all_plate_data:
                messagebox.showinfo("ℹ️ Bilgi", "Kaydedilecek veri bulunamadı.")
                return
                
            # Veritabanı bağlantısı
            db_path = os.path.join(os.getcwd(), "plates_database.db")
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            
            # Tablo oluştur (eğer yoksa)
            cursor.execute('''
            CREATE TABLE IF NOT EXISTS plates (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                tarih_saat TEXT,
                plaka TEXT,
                dogruluk REAL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
            ''')
            
            # Her kaydı kontrol et ve ekle
            kayit_sayisi = 0
            for plate in self.all_plate_data:
                # Önce bu kaydın veritabanında olup olmadığını kontrol et
                cursor.execute('''
                SELECT COUNT(*) FROM plates 
                WHERE tarih_saat = ? AND plaka = ?
                ''', (plate['tarih_saat'], plate['plaka']))
                
                if cursor.fetchone()[0] == 0:  # Eğer kayıt yoksa
                    # Doğruluk oranını sayısal değere çevir
                    dogruluk_str = plate['dogruluk'].replace('%', '')
                    dogruluk_float = float(dogruluk_str) / 100.0
                    
                    cursor.execute('''
                    INSERT INTO plates (tarih_saat, plaka, dogruluk)
                    VALUES (?, ?, ?)
                    ''', (plate['tarih_saat'], plate['plaka'], dogruluk_float))
                    kayit_sayisi += 1
            
            # Değişiklikleri kaydet ve bağlantıyı kapat
            conn.commit()
            conn.close()
            
            if kayit_sayisi > 0:
                messagebox.showinfo("✅ Başarılı", f"💾 {kayit_sayisi} yeni kayıt veritabanına eklendi!\nVeritabanı: {db_path}")
            else:
                messagebox.showinfo("ℹ️ Bilgi", "Tüm kayıtlar zaten veritabanında mevcut.")
        except Exception as e:
            messagebox.showerror("❌ Hata", f"Veritabanına kaydetme hatası: {e}")
    
    def on_closing(self):
        """Uygulama kapatıldığında çalışan işlemi sonlandır"""
        try:
            self.terminate_process()
        except Exception as e:
            print(f"Kapanış sırasında hata: {e}")
        finally:
            self.root.destroy()

def main():
    """Ana fonksiyon - PyInstaller uyumluluğu için"""
    try:
        # Tkinter uygulamasını başlat
        root = tk.Tk()
        
        # Windows'ta taskbar ikonunu ayarla
        try:
            root.iconbitmap(default=get_resource_path("icon.ico"))
        except:
            pass  # İkon yoksa geç
        
        app = PlateRecognitionApp(root)
        
        # Ana döngüyü başlat
        root.mainloop()
        
    except Exception as e:
        # Hata durumunda kullanıcıya mesaj göster
        try:
            import tkinter.messagebox as mb
            mb.showerror("Kritik Hata", f"Uygulama başlatılırken hata oluştu:\n{e}")
        except:
            print(f"Kritik Hata: {e}")
            input("Devam etmek için Enter'a basın...")

if __name__ == "__main__":
    main()